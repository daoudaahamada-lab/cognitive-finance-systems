"""Lit le classeur source (onglets Planning et Slides) et écrit un JSON par carrousel."""
import argparse
import json
import unicodedata
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "data" / "WTH_Octobre_V2_Textes_Slides.xlsx"
OUT = ROOT / "data" / "carousels"


def clean(value):
    if value is None:
        return ""
    return unicodedata.normalize("NFC", str(value)).strip()


def load(xlsx=XLSX):
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    carousels = {}

    rows = wb["Planning"].iter_rows(values_only=True)
    header = [clean(h) for h in next(rows)]
    for row in rows:
        rec = dict(zip(header, row))
        cid = clean(rec.get("ID"))
        if not cid:
            continue
        carousels[cid] = {
            "id": cid,
            "famille": clean(rec.get("Famille")),
            "format": clean(rec.get("Format")),
            "accroche": clean(rec.get("Accroche (slide 1)")),
            "sujet": clean(rec.get("Sujet d’origine")),
            "legende": clean(rec.get("Légende")),
            "hashtags": clean(rec.get("Hashtags")),
            "slides": [],
        }

    rows = wb["Slides"].iter_rows(values_only=True)
    header = [clean(h) for h in next(rows)]
    for row in rows:
        rec = dict(zip(header, row))
        cid = clean(rec.get("ID"))
        if cid not in carousels:
            continue
        carousels[cid]["slides"].append({
            "n": int(clean(rec.get("N° slide")).split("/")[0]),
            "fichier": clean(rec.get("Fichier attendu")),
            "titre": clean(rec.get("Titre")),
            "texte": clean(rec.get("Texte")),
        })

    for c in carousels.values():
        c["slides"].sort(key=lambda s: s["n"])
    return carousels


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("ids", nargs="*", help="IDs carrousel (WTH-003…). Vide = tous.")
    args = parser.parse_args()
    carousels = load()
    OUT.mkdir(parents=True, exist_ok=True)
    for cid in args.ids or sorted(carousels):
        c = carousels[cid]
        if len(c["slides"]) != 9:
            raise SystemExit(f"{cid}: {len(c['slides'])} slides au lieu de 9")
        (OUT / f"{cid}.json").write_text(json.dumps(c, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(f"{cid}: 9 slides -> {OUT / (cid + '.json')}")


if __name__ == "__main__":
    main()
